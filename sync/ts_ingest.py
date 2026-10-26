# Usage: python3 -I ts_ingest.py OUTDIR RAWJSON [RAWJSON ...]
# RAWJSON = saved result of Gmail get_message(messageFormat=RAW). Writes one snapshot JSON per report into OUTDIR and prints doc ids.
import sys,json,base64,email,io,re,os,datetime,warnings
warnings.filterwarnings("ignore")
import openpyxl
def dept(*vals):
    for v in vals:
        u=str(v or "").lower()
        if not u: continue
        if re.search(r"sewer|drain",u): return "Sewer"
        if "plumb" in u: return "Plumbing"
        if "elec" in u: return "Electrical"
        if re.search(r"hvac|heat|cool|air|furnace",u): return "HVAC"
    return "Other"
def pdate(s):
    m=re.match(r"(\d{1,2})/(\d{1,2})/(\d{2,4})",s.strip()); y=int(m.group(3)); y+=2000 if y<100 else 0
    return datetime.date(y,int(m.group(1)),int(m.group(2))).isoformat()
def mins(v):
    if v in (None,""): return 0
    if isinstance(v,(int,float)): return round(v*60)
    try: return round(float(str(v).replace(",",""))*60)
    except: return 0
def parse(xbytes,captured_iso,source):
    wb=openpyxl.load_workbook(io.BytesIO(xbytes),data_only=True)
    ws=wb.worksheets[0]; rows=list(ws.iter_rows(values_only=True))
    hi=None
    for i,r in enumerate(rows[:15]):
        h=[str(x or "").strip().lower() for x in r]
        hrs=next((j for j,t in enumerate(h) if re.search(r"total paid time|total hours worked",t)),None)
        nm=next((j for j,t in enumerate(h) if re.fullmatch(r"name|technician( name)?|employee( name)?",t)),None)
        if hrs is not None and nm is not None:
            hi=i; c=dict(name=nm,hours=hrs,trade=next((j for j,t in enumerate(h) if "trade" in t),-1),
                         team=h.index("team") if "team" in h else -1,unit=next((j for j,t in enumerate(h) if t.endswith("business unit")),-1)); break
    if hi is None: raise SystemExit("no header row in "+source)
    emps=[]
    for r in rows[hi+1:]:
        name=str(r[c["name"]] or "").strip()
        if not name or re.fullmatch(r"\d+",name) or re.match(r"(office|warehouse|parking lot)",name,re.I): continue
        g=lambda j: str(r[j] or "").strip() if j>=0 else ""
        emps.append(dict(name=name,unit=g(c["unit"]),team=g(c["team"]),dept=dept(g(c["trade"]),g(c["team"]),g(c["unit"])),min=mins(r[c["hours"]])))
    rs=re_=None
    for w in wb.worksheets:
        for r in w.iter_rows(values_only=True):
            if r and re.search(r"date range|from - to",str(r[0] or ""),re.I) and len(r)>1 and r[1]:
                p=re.split(r"\s+[-–]\s+",str(r[1]).strip()); rs=pdate(p[0]); re_=pdate(p[-1])
    if not rs: rs=re_=captured_iso[:10]
    return dict(captured_at=captured_iso,range_start=rs,range_end=re_,source=source,employees=emps)
out=sys.argv[1]; os.makedirs(out,exist_ok=True)
for f in sys.argv[2:]:
    d=json.load(open(f)); raw=d["raw"]
    m=email.message_from_bytes(base64.urlsafe_b64decode(raw+"="*(-len(raw)%4)))
    # report "as of" time, rounded to the minute so duplicate copies of one email collapse
    t=datetime.datetime.fromtimestamp(int(d["internalDate"])/1000,datetime.timezone.utc).replace(second=0,microsecond=0)
    cap=t.strftime("%Y-%m-%dT%H:%M:00.000Z")
    for p in m.walk():
        fn=p.get_filename()
        if fn and fn.lower().endswith(".xlsx"):
            s=parse(p.get_payload(decode=True),cap,fn)
            doc=s["range_start"]+"_"+s["range_end"]+"_"+cap.replace(":","-").replace(".","-")
            json.dump(s,open(os.path.join(out,doc+".json"),"w"))
            print(doc, len([e for e in s["employees"] if e["min"]>0]), "with hours")
